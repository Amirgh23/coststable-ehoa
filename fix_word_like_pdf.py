from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(r'C:\Users\a.ghafarian\Downloads\SUBMIT_TO_PROFESSOR_MOATTAR')
src = root / 'CostStableEHOA_English_BAIC2026_Compact.docx'
out = root / 'CostStableEHOA_English_BAIC2026_Native.docx'
doc = Document(src)

# PDF reading order: Method heading and Equation (1), then the repair algorithm.
paras = doc.paragraphs
algorithm = paras[16:22]
anchor = paras[26]._p
for paragraph in algorithm:
    anchor.addnext(paragraph._p)
    anchor = paragraph._p

# Keep every table row intact and prevent captions/headings from becoming orphans.
for table in doc.tables:
    for row_index, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if row_index < len(table.rows) - 1:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
for p in doc.paragraphs:
    text = p.text.strip()
    if p.style.name.startswith('Heading') or text.startswith(('TABLE I', 'TABLE II', 'Fig.')):
        p.paragraph_format.keep_with_next = True

paras = doc.paragraphs
for i, p in enumerate(paras):
    text = p.text.strip()
    if text == 'Stable-LCB repair':
        p.paragraph_format.keep_with_next = True
    if text in ('TABLE I', 'TABLE II'):
        j = i
        while j + 1 < len(paras) and paras[j + 1].text.strip() and not paras[j + 1].text.strip().startswith('Fig.'):
            paras[j].paragraph_format.keep_with_next = True
            j += 1

doc.core_properties.title = 'CostStableEHOA Budget Feasible and Stability Aware Feature Selection'
doc.save(out)
print(out)
