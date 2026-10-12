from pathlib import Path
import csv, json, re, shutil
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'revision_2026_10_12'
TEMPLATE=Path(r'C:\Users\a.ghafarian\Downloads\قالب_مقاله.docx')
if not TEMPLATE.exists():TEMPLATE=OUT/'sources/congress_template.docx'
def read(name):
    with (OUT/'analysis'/name).open(encoding='utf-8',newline='') as f: return list(csv.DictReader(f))
ds=read('datasets.csv'); overall=read('eligible_overall.csv'); fr=read('eligible_friedman.csv'); ranks=read('eligible_ranks.csv'); pairs=read('eligible_paired.csv'); bc=read('breast_cancer_case.csv')
methods=['EHOA','Budget-random','Budget-cheapest','Budget-mi','Budget-stable_lcb']
labels=['EHOA','Random/cost','Cheapest','MI/cost','Resamp./cost']
names=dict(breast_cancer='Breast Cancer',wine='Wine',ionosphere='Ionosphere',sonar='Sonar',dermatology='Dermatology',spambase='Spambase',parkinsons='Parkinsons*',hill_valley='Hill-Valley',steel_plates_fault='Steel Plates*',madelon='Madelon')
references=[
('R. Jagdhuber, M. Lang, A. Stenzl, J. Neuhaus, and J. Rahnenführer, "Cost-Constrained feature selection in binary classification: adaptations for greedy forward selection and genetic algorithms," BMC Bioinformatics, vol. 21, art. 26, 2020.','https://doi.org/10.1186/s12859-020-3361-9'),
('S. Nogueira, K. Sechidis, and G. Brown, "On the Stability of Feature Selection Algorithms," Journal of Machine Learning Research, vol. 18, no. 174, pp. 1-54, 2018.','https://jmlr.org/papers/v18/17-514.html'),
('A. E. Hegazy, M. A. Makhlouf, O. A. M. Salem, and B. Hafiz, "An enhanced Hiking optimization algorithm for accurate and interpretable feature selection in medical data classification," Cluster Computing, vol. 29, art. 244, 2026.','https://doi.org/10.1007/s10586-026-05946-9'),
('V. Hamer and P. Dupont, "An Importance Weighted Feature Selection Stability Measure," Journal of Machine Learning Research, vol. 22, no. 116, pp. 1-57, 2021.','https://jmlr.org/papers/v22/20-366.html'),
('T. Klonecki and P. Teisseyre, "Feature selection under budget constraint in medical applications: analysis of penalized empirical risk minimization methods," Applied Intelligence, vol. 53, pp. 29943-29973, 2023.','https://doi.org/10.1007/s10489-023-05063-9'),
('M. Valancius, M. Lennon, and J. Oliva, "Acquisition Conditioned Oracle for Nongreedy Active Feature Acquisition," Proceedings of the 41st International Conference on Machine Learning, PMLR, vol. 235, pp. 48957-48975, 2024.','https://proceedings.mlr.press/v235/valancius24a.html'),
('T. Klonecki, P. Teisseyre, and J. Lee, "Cost-Constrained Multi-label Group Feature Selection Using Shadow Features," in Advances in Artificial Intelligence Research, PP-RAI 2025, Lecture Notes in Networks and Systems, vol. 1599, pp. 3-16, Springer, 2026.','https://doi.org/10.1007/978-3-032-04197-5_1'),
('J. Demšar, "Statistical Comparisons of Classifiers over Multiple Data Sets," Journal of Machine Learning Research, vol. 7, pp. 1-30, 2006.','https://jmlr.org/papers/v7/demsar06a.html'),
('W. Wolberg, O. Mangasarian, N. Street, and W. Street, "Breast Cancer Wisconsin (Diagnostic)," UCI Machine Learning Repository, 1993.','https://doi.org/10.24432/C5DW2B')]
references.extend([
('B. Remeseiro and V. Bolon-Canedo, "A review of feature selection methods in medical applications," Computers in Biology and Medicine, vol. 112, art. 103375, 2019.','https://doi.org/10.1016/j.compbiomed.2019.103375'),
('N. Meinshausen and P. Buhlmann, "Stability selection," Journal of the Royal Statistical Society: Series B, vol. 72, no. 4, pp. 417-473, 2010.','https://doi.org/10.1111/j.1467-9868.2010.00740.x'),
('R. D. Shah and R. J. Samworth, "Variable selection with error control: another look at stability selection," Journal of the Royal Statistical Society: Series B, vol. 75, no. 1, pp. 55-80, 2013.','https://doi.org/10.1111/j.1467-9868.2011.01034.x'),
('B. C. Ross, "Mutual Information between Discrete and Continuous Data Sets," PLOS ONE, vol. 9, no. 2, art. e87357, 2014.','https://doi.org/10.1371/journal.pone.0087357'),
('N. V. Chawla, K. W. Bowyer, L. O. Hall, and W. P. Kegelmeyer, "SMOTE: Synthetic Minority Over-sampling Technique," Journal of Artificial Intelligence Research, vol. 16, pp. 321-357, 2002.','https://doi.org/10.1613/jair.953'),
('S. Varma and R. Simon, "Bias in error estimation when using cross-validation for model selection," BMC Bioinformatics, vol. 7, art. 91, 2006.','https://doi.org/10.1186/1471-2105-7-91')])
(OUT/'references_for_submission.txt').write_text('\n'.join(t+' '+u for t,u in references)+'\n',encoding='utf-8')
(OUT/'sources/reference_verification.csv').write_text('reference_id,verified_source,verification_date,doi_status\n'+'\n'.join(f'{i},"{u}",2026-10-12,"'+('Publisher DOI verified' if 'doi.org' in u else 'No DOI assigned on publisher page; direct publisher URL')+'"' for i,(t,u) in enumerate(references,1))+'\n',encoding='utf-8')
shutil.copy2(OUT/'sources/image1.png',OUT/'baic2026_header.png')

def h(en,fa): return dict(kind='h',en=en,fa=fa)
def p(en,fa): return dict(kind='p',en=en,fa=fa)
def eq(tex,word): return dict(kind='eq',tex=tex,word=word)
def table(cap_en,cap_fa,headers,rows,widths=None): return dict(kind='table',en=cap_en,fa=cap_fa,headers=headers,rows=rows,widths=widths)

title={'en':'Hard-Budget Feature Selection with Resampled Mutual-Information Repair in EHOA','fa':'انتخاب ویژگی با بودجه سخت و تعمیر مبتنی بر اطلاعات متقابل بازنمونه‌گیری‌شده در EHOA'}
abstract={'en':'We compare four budget repairs within the Enhanced Hiking Optimization Algorithm (EHOA) while preserving its search and wrapper objective. The proposed repair ranks features by a nonnegative mean-minus-standard-deviation mutual-information score divided by synthetic acquisition cost. This score is a conservative heuristic, not a calibrated confidence bound. We audit 300 archived runs on ten classification datasets. All 240 constrained outputs satisfy a budget of 30% of total feature cost. Source inspection identifies target leakage in Steel Plates and a subject-overlap risk in Parkinsons; we exclude both from the main retrospective analysis. Across eight eligible datasets, resampled repair changes mean balanced accuracy from 0.8068 to 0.8116, cost ratio from 0.4791 to 0.2792, and Nogueira stability from 0.1161 to 0.2550. Friedman tests detect method differences in cost and stability, but not balanced accuracy. The resampled-versus-EHOA stability contrast does not survive within-metric Holm correction. A breast-cancer cytology case illustrates subset selection without claiming clinical validation or monetary savings. The contribution is an auditable repair comparison, not general predictive superiority.',
'fa':'چهار روش تعمیر بودجه را در الگوریتم بهینه‌سازی کوه‌پیمایی بهبود‌یافته (EHOA) با جست‌وجو و تابع هدف یکسان مقایسه می‌کنیم. روش پیشنهادی، اطلاعات متقابل را بازنمونه‌گیری می‌کند و نسبت امتیاز نامنفی «میانگین منهای انحراف معیار» به هزینه مصنوعی را مبنای حذف قرار می‌دهد. این امتیاز ابتکاری محافظه‌کارانه است، نه کران اطمینان کالیبره‌شده. ممیزی ۳۰۰ اجرای موجود روی ده دیتاست نشان می‌دهد همه ۲۴۰ خروجی مقید، بودجه ۳۰ درصد هزینه کل را رعایت می‌کنند. به‌دلیل نشت هدف در Steel Plates و خطر هم‌پوشانی افراد در Parkinsons، تحلیل اصلی گذشته‌نگر به هشت دیتاست محدود شد. تعمیر بازنمونه‌گیری‌شده، میانگین دقت متوازن را از 0.8068 به 0.8116، نسبت هزینه را از 0.4791 به 0.2792 و پایداری نوگوئیرا را از 0.1161 به 0.2550 تغییر داد. فریدمن اختلاف روش‌ها را برای هزینه و پایداری نشان داد، اما برای دقت متوازن معنادار نبود. اختلاف زوجی پایداری با EHOA پس از اصلاح هولم درون معیار معنادار نیست. مطالعه موردی سیتولوژی سرطان پستان، کاربرد انتخاب زیرمجموعه را نشان می‌دهد و اعتبارسنجی بالینی یا صرفه‌جویی پولی را ادعا نمی‌کند. سهم پژوهش، مقایسه قابل‌ممیزی تعمیرهاست، نه برتری عمومی پیش‌بینی.'}
keywords={'en':'feature selection, hard budget, resampled mutual information, subset stability','fa':'انتخاب ویژگی، بودجه سخت، اطلاعات متقابل بازنمونه‌گیری‌شده، پایداری زیرمجموعه'}
abstract['en']+=' A separate 60-run exploratory correction also finds no general predictive advantage.'
abstract['fa']+=' شصت اجرای اصلاح‌شده اکتشافی جداگانه نیز برتری عمومی پیش‌بینی را نشان نداد.'

related_rows=[['Jagdhuber [1]','Yes','Additive','NR','Simulation','Feasible GA'],['Nogueira [2]','NR','NR','Chance-corrected','Resampling','Metric'],['Hamer [4]','NR','NR','Importance-weighted','Resampling','Metric limits'],['Klonecki [5]','Yes','Feature / proxy','NR','Medical data','Penalized ERM'],['Valancius [6]','NR','Sequential','NR','AFA benchmarks','Patient-wise policy'],['Klonecki [7]','Yes','Grouped','NR','MIMIC','Shadow stopping'],['This work','Yes','Synthetic additive','Nogueira','2 x 3 outer','MI/cost repair']]
page1=[
h('1  Problem and contribution','۱  مسئله و سهم پژوهش'),
p('A small feature set need not be inexpensive: acquisition costs can differ between variables. Hard-budget selection limits the sum of selected costs [1]. Subset stability measures how feature choices vary across training samples and random search, rather than how accuracy or convergence varies [2]. In biomedical classification, this distinction matters when investigators interpret a selected measurement panel.', 'زیرمجموعه کوچک الزاماً کم‌هزینه نیست؛ هزینه کسب ویژگی‌ها می‌تواند متفاوت باشد. انتخاب با بودجه سخت، مجموع هزینه ویژگی‌های منتخب را محدود می‌کند [1]. پایداری زیرمجموعه، تغییر انتخاب ویژگی در نمونه‌های آموزش و جست‌وجوی تصادفی را می‌سنجد؛ این مفهوم با ثبات دقت یا همگرایی متفاوت است [2]. این تمایز برای تفسیر مجموعه اندازه‌گیری‌های زیست‌پزشکی اهمیت دارد.'),
p('We retain the implemented EHOA search [3] and compare unconstrained EHOA with four repairs. Only the repair utility changes among constrained variants. The contribution is the integration and controlled evaluation of resampled relevance-per-cost repair. We do not claim that budgets, repair, mutual information (MI), or resampling are new. CostStableEHOA remains the code identifier; the revised title avoids implying stable costs or proven selection stability.', 'جست‌وجوی پیاده‌سازی‌شده EHOA [3] را حفظ و آن را با چهار تعمیر مقایسه می‌کنیم. بین نسخه‌های مقید، فقط سودمندی تعمیر تغییر می‌کند. سهم پژوهش، ادغام و ارزیابی کنترل‌شده تعمیر مبتنی بر ارتباط بازنمونه‌گیری‌شده به‌ازای هزینه است. بودجه، تعمیر، اطلاعات متقابل (MI) و بازنمونه‌گیری به‌تنهایی نوآوری نیستند. شناسه کد CostStableEHOA حفظ شد؛ عنوان جدید، ثبات هزینه یا تضمین پایداری انتخاب را القا نمی‌کند.'),
h('2  Related studies','۲  مطالعات مرتبط'),
p('Stability metrics can reward reproducible yet weak predictors; importance-weighted stability addresses this concern [4]. Medical budget selection with penalized risk [5], sequential acquisition [6], and grouped costs with shadow features [7] solve related but different problems. Table 1 summarizes these distinctions; NR means not reported in the inspected source. Their numerical results are not directly comparable with our protocol.', 'شاخص پایداری ممکن است انتخاب‌های تکرارپذیر اما کم‌ارزش را مطلوب نشان دهد؛ پایداری وزن‌دار با اهمیت ویژگی به این مسئله می‌پردازد [4]. انتخاب پزشکی با ریسک جریمه‌شده [5]، کسب ترتیبی ویژگی [6] و هزینه گروهی با ویژگی سایه [7] مسائل مرتبط ولی متفاوتی‌اند. جدول ۱ تفاوت‌ها را نشان می‌دهد. NR یعنی گزارش نشده در منبع بررسی‌شده؛ نتایج عددی این مطالعات با پروتکل حاضر مستقیماً مقایسه نمی‌شوند.'),
table('Table 1. Related-work comparison. Full evaluation details are in the supplement.','جدول ۱. مقایسه مطالعات؛ جزئیات پروتکل در پیوست آمده است.', ['Study','Budget','Costs','Stability','Evaluation','Difference'],related_rows),
h('3  Implemented EHOA','۳  EHOA پیاده‌سازی‌شده'),
p('Each hiker has a continuous position x and a velocity v; a binary mask z selects features. Initialization uses a tent sequence with initial value 0.7 and parameter 0.4, mapped to [-6,6]. Velocities start uniformly in [-1.2,1.2]. For iteration t of T=10, the implemented update is:', 'هر عضو جمعیت، موقعیت پیوسته x و سرعت v دارد؛ ماسک دودویی z ویژگی‌ها را انتخاب می‌کند. مقداردهی اولیه از دنباله tent با مقدار آغازین 0.7 و پارامتر 0.4 در بازه [-6,6] استفاده می‌کند. سرعت اولیه در [-1.2,1.2] یکنواخت است. به‌روزرسانی پیاده‌سازی در تکرار t از T=10 چنین است:'),
eq(r'v_i\leftarrow\operatorname{clip}_{[-12,12]}\!\left(w_t v_i+SF_t[2r_1(p_i-x_i)+2r_2(g-x_i)]\right)', 'vᵢ ← clip[−12,12](wₜvᵢ + SFₜ[2r₁(pᵢ−xᵢ)+2r₂(g−xᵢ)])'),
eq(r'x_i\leftarrow\operatorname{clip}_{[-6,6]}(x_i+v_i),\quad w_t=0.9-0.5t/T,\quad SF_t=3-2t/T', 'xᵢ ← clip[−6,6](xᵢ+vᵢ); wₜ = 0.9−0.5t/T; SFₜ = 3−2t/T'),
p('Here p_i and g are personal and global best continuous positions, and r_1,r_2 are independent per-feature uniform draws. The sigmoid gives selection probability 1/(1+exp(-x_j)); an empty sampled mask restores the largest probability. The implementation keeps continuous leaders separate from repaired masks, caches fitness for duplicate masks, and stops after ten iterations without early stopping.', 'در این روابط p_i و g بهترین موقعیت پیوسته فردی و سراسری‌اند؛ r_1 و r_2 برای هر ویژگی مستقل و یکنواخت‌اند. احتمال انتخاب از سیگموئید 1/(1+exp(-x_j)) می‌آید. ماسک نمونه‌گیری‌شده تهی، ویژگی دارای بیشترین احتمال را بازمی‌گرداند. کد، رهبرهای پیوسته را از ماسک تعمیرشده جدا نگه می‌دارد، برازندگی ماسک تکراری را ذخیره می‌کند و پس از ده تکرار بدون توقف زودهنگام پایان می‌یابد.')]

dataset_rows=[]
for row in ds:
    counts='/'.join(str(v) for v in json.loads(row['class_counts_used']).values())
    src='U'+str(dict(breast_cancer=17,wine=109).get(row['dataset'],'')) if row['dataset'] in ['breast_cancer','wine'] else 'O'+str(dict(ionosphere=59,sonar=40,dermatology=35,spambase=44,parkinsons=1488,hill_valley=1479,steel_plates_fault=1504,madelon=1485)[row['dataset']])
    dataset_rows.append([names[row['dataset']],f"{row['N_snapshot']} / {row['N_used']}",row['D_used'],row['classes'],counts,src])
page2=[h('4  Budget repair and stability','۴  تعمیر بودجه و پایداری'),
p('Costs are fixed positive additive synthetic units, independent of labels: c is drawn once per dataset from lognormal(0,0.7), with NumPy seed 1000+D. The absolute budget B=0.30 sum(c) is derived from a relative budget. Ionosphere and Dermatology share a cost vector because both have D=34. A nonempty mask is feasible when sum(z_j c_j) <= B.', 'هزینه‌ها، واحدهای مصنوعی مثبت، جمع‌پذیر و مستقل از برچسب‌اند. برای هر دیتاست یک‌بار از lognormal(0,0.7) با بذر NumPy برابر 1000+D تولید می‌شوند. بودجه مطلق B=0.30 sum(c) از نسبت بودجه حاصل می‌شود. Ionosphere و Dermatology به‌دلیل D=34، بردار هزینه یکسان دارند. ماسک غیرتهی زمانی شدنی است که sum(z_j c_j) <= B باشد.'),
eq(r'u_j=\max(0,\bar I_j-\lambda s_j),\qquad q_j=u_j/c_j,\qquad\lambda=1', 'uⱼ = max(0, Īⱼ−λsⱼ); qⱼ = uⱼ/cⱼ; λ = 1'),
p('For resampled repair, four within-class samples draw ceil(0.8 n_class) rows with replacement (minimum two). Dense mutual_info_classif uses continuous-feature estimation, three neighbors, and nats; missing values are zero-filled for MI only. s_j is the sample standard deviation (ddof=1). No division by sqrt(R) or calibrated confidence level occurs. Utility is computed once on the outer-training partition using a separate RNG seeded by selector_seed+7919.', 'در تعمیر بازنمونه‌گیری‌شده، چهار نمونه درون هر کلاس با جایگذاری و اندازه ceil(0.8 n_class)، حداقل دو، تولید می‌شود. mutual_info_classif برای ماتریس متراکم، ویژگی‌ها را پیوسته و تعداد همسایه را سه در نظر می‌گیرد؛ واحد MI نَت است. فقط برای MI، مقدار گم‌شده صفر می‌شود. s_j انحراف معیار نمونه با ddof=1 است. تقسیم بر sqrt(R) یا سطح اطمینان کالیبره‌شده وجود ندارد. سودمندی یک‌بار روی آموزش بیرونی و با بذر مستقل selector_seed+7919 محاسبه می‌شود.'),
table('Table 2. Four repairs share the same deletion rule.','جدول ۲. هر چهار تعمیر، قاعده حذف یکسان دارند.', ['Repair','Utility u_j','Meaning'],[['Random/cost','U(0,1), fixed per fit','Not uniform deletion'],['Cheapest','1','Largest cost first'],['MI/cost','Single-sample MI','Relevance per cost'],['Resampled/cost','max(0, mean MI - SD)','Conservative heuristic']]),
p('Procedure: initialize population; sample sigmoid mask; repair before fitness; evaluate; update personal/global leaders; repeat movement, sampling, repair, and evaluation; return the best repaired mask. Repair removes argmin(q_j) among active features while cost exceeds B and more than one feature remains. An unaffordable singleton is cleared; an empty mask restores the affordable feature maximizing q_j. A tie selects the first feature index. Utility is never recomputed per candidate.', 'روند اجرا: جمعیت را مقداردهی کن؛ ماسک سیگموئید را بساز؛ پیش از برازندگی تعمیر کن؛ ارزیابی و رهبرهای فردی/سراسری را به‌روزرسانی کن؛ حرکت، نمونه‌گیری، تعمیر و ارزیابی را تکرار کن؛ بهترین ماسک تعمیرشده را برگردان. تعمیر تا زمان نقض بودجه و وجود بیش از یک ویژگی، کمترین q_j فعال را حذف می‌کند. تک‌ویژگی غیرقابل‌خرید حذف و ماسک تهی با بهترین نسبت از بین ویژگی‌های قابل‌خرید پر می‌شود. تساوی با اولین اندیس شکسته می‌شود. سودمندی برای هر جواب دوباره محاسبه نمی‌شود.'),
eq(r'F(z)=0.99(1-\overline{Acc}_{inner})+0.01\|z\|_0/D', 'F(z) = 0.99(1−Acc̄inner) + 0.01|z|/D'),
p('The unchanged objective minimizes ordinary inner accuracy error, not balanced-accuracy error. External balanced accuracy (BA) is the mean recall across classes. Stability is computed once per dataset/method from M=6 final masks:', 'تابع هدف دست‌نخورده، خطای دقت معمولی درونی را کمینه می‌کند، نه خطای دقت متوازن. دقت متوازن بیرونی (BA)، میانگین بازخوانی کلاس‌هاست. پایداری یک‌بار برای هر دیتاست/روش و از M=6 ماسک نهایی محاسبه می‌شود:'),
eq(r'\Phi=\operatorname{clip}_{[-1,1]}\left(1-\frac{\frac{M}{M-1}\frac{1}{D}\sum_j p_j(1-p_j)}{(\bar k/D)(1-\bar k/D)}\right)', 'Φ = clip[−1,1](1 − {M/(M−1) · meanⱼ[pⱼ(1−pⱼ)]}/{(k̄/D)(1−k̄/D)})'),
p('p_j is selection frequency and k-bar is mean subset size. Identical empty/full masks return one in the code. This dimensionless chance-corrected outcome [2] mixes data-split and search variation; overlapping training folds make the six masks dependent. We do not treat copied stability values in six CSV rows as six observations. Jaccard overlap is supplementary and is not chance-corrected.', 'p_j فراوانی انتخاب و k-bar میانگین اندازه زیرمجموعه است. کد برای ماسک‌های تهی/کامل یکسان، مقدار یک برمی‌گرداند. این شاخص بدون واحد و تصحیح‌شده نسبت به شانس [2]، تغییر ناشی از تقسیم داده و جست‌وجو را با هم می‌سنجد. هم‌پوشانی آموزش، شش ماسک را وابسته می‌کند. شش مقدار پایداری تکرارشده در CSV شش مشاهده مستقل نیستند. هم‌پوشانی جاکارد، معیار تکمیلی بدون تصحیح شانس است.'),
h('5  Datasets and source audit','۵  دیتاست‌ها و ممیزی منبع'),
table('Table 3. Snapshot / deduplicated N; D is the supplied numeric feature count.','جدول ۳. N فایل / پس از حذف تکراری؛ D تعداد ستون عددی موجود است.', ['Dataset','N raw/used','D','K','Class counts used','Source'],dataset_rows),
p('U identifiers link to archive.ics.uci.edu/dataset/{id}; O identifiers link to openml.org/d/{id}. The supplement gives clickable direct links, raw/used class counts, hashes and budgets. No runtime one-hot encoding occurs. Steel has 33 inputs and two classes in the snapshot, whereas UCI defines 27 predictors and seven fault types. Its V28-V33 are six fault indicators and exactly determine the target; we exclude these results. Parkinsons is a recording-level dataset with repeated participants and no subject-group split; we also exclude it from the main analysis. Eight eligible datasets retain 240 runs.', 'شناسه U به archive.ics.uci.edu/dataset/{id} و O به openml.org/d/{id} ارجاع می‌دهد. پیوست، لینک مستقیم، توزیع خام/پاک‌شده، هش و بودجه را دارد. کدگذاری one-hot هنگام اجرا انجام نمی‌شود. Steel در فایل ۳۳ ورودی و دو کلاس دارد، ولی UCI آن را با ۲۷ پیش‌بین و هفت نوع عیب معرفی می‌کند. ستون‌های V28 تا V33 شش نشانگر عیب‌اند و هدف را دقیق تعیین می‌کنند؛ نتایج آن کنار گذاشته شد. Parkinsons شامل ضبط‌های تکراری افراد است و تقسیم فردمحور ندارد؛ آن نیز از تحلیل اصلی حذف شد. هشت دیتاست باقی‌مانده ۲۴۰ اجرا دارند.')]

param_rows=[['Population / iterations','8 / 10; no early stop'],['Position / velocity clip','[-6,6] / [-12,12]'],['Tent / sweep / inertia','0.7, p=0.4 / 3 to 1 / 0.9 to 0.4'],['c1,c2 / alpha','2,2 / 0.99'],['Outer / inner folds','3 stratified / 3 stratified'],['Outer seeds','71001, 71002'],['Selector seed','outer seed + fold index x 100003'],['MI resamples / fraction','4 / 0.8 with replacement'],['lambda / MI neighbors','1 / 3'],['Cost / budget','lognormal(0,0.7) / 30% total'],['Classifier','5-NN; uniform; Euclidean'],['Imputation / scale','Training mean / StandardScaler'],['SMOTE','Training only; k=min(5,n_min-1)']]
mainrows=[]
for m,label in zip(methods,labels):
    r=next(x for x in overall if x['method']==m)
    mainrows.append([label]+[f"{float(r[x]):.4f}" if x!='selected_features' else f"{float(r[x]):.2f}" for x in ['balanced_accuracy','cost_ratio','selected_features','nogueira_stability']]+['47/48' if m=='EHOA' else '0/48'])
frows=[]
metriclabels={'balanced_accuracy':'BA','cost_ratio':'Cost','selected_features':'Size','nogueira_stability':'Stability'}
for metric,label in metriclabels.items():
    r=next(x for x in fr if x['metric']==metric)
    frows.append([label,f"{float(r['statistic']):.4f}",f"{float(r['p_raw']):.5f}"]+[f"{float(next(x for x in ranks if x['metric']==metric and x['method']==m)['average_rank']):.2f}" for m in methods])
page3=[h('6  Evaluation and statistical analysis','۶  ارزیابی و تحلیل آماری'),
p('Table 4 documents actual settings rather than class defaults. All methods share data splits, costs, seeds, search equations and wrapper fitness. Initialization positions are deterministic; search velocities and masks use the selector RNG. The four repairs therefore isolate utility ranking under a common budget; unconstrained EHOA additionally tests the effect of imposing that budget.', 'جدول ۴ تنظیمات واقعی اجرا را نشان می‌دهد، نه پیش‌فرض کلاس. تقسیم‌ها، هزینه‌ها، بذرها، معادلات جست‌وجو و برازندگی برای همه روش‌ها یکسان‌اند. موقعیت اولیه قطعی است؛ سرعت و ماسک از مولد تصادفی انتخابگر می‌آیند. چهار تعمیر، اثر رتبه‌بندی سودمندی را تحت بودجه مشترک تفکیک می‌کنند؛ مقایسه با EHOA همچنین اثر اعمال بودجه را می‌سنجد.'),
table('Table 4. Settings extracted from code and configuration.','جدول ۴. پارامترهای استخراج‌شده از کد و پیکربندی.', ['Parameter','Value'],param_rows),
p('Within each inner fold, training data alone fit mean imputation, scaling and SMOTE; validation rows remain unchanged. MI utility uses all outer-training labels, including rows that later become inner validation. Thus inner scores are selection scores, not unbiased estimates. Outer-test labels do not enter selection. Final outer training fits preprocessing/SMOTE and 5-NN; archived metrics retain only 5-NN although the runner fits four classifiers. Preprocessing and SMOTE occur before masking; this is a simulation, not a deployed selected-only acquisition pipeline.', 'در هر تقسیم درونی، فقط آموزش برای جایگذاری میانگین، استانداردسازی و SMOTE به‌کار می‌رود؛ اعتبارسنجی بازنمونه‌گیری نمی‌شود. سودمندی MI تمام برچسب‌های آموزش بیرونی، از جمله ردیف‌های اعتبارسنجی درونی بعدی، را می‌بیند؛ بنابراین امتیاز درونی، امتیاز انتخاب است و برآورد بی‌اریب نیست. برچسب آزمون بیرونی وارد انتخاب نمی‌شود. آموزش نهایی، پیش‌پردازش/SMOTE و 5-NN را برازش می‌کند. اجراگر چهار طبقه‌بند را برازش کرده، اما آرشیو فقط معیارهای 5-NN را نگه داشته است. پیش‌پردازش و SMOTE پیش از اعمال ماسک‌اند؛ این شبیه‌سازی است، نه سامانه عملیاتی کسب فقط ویژگی‌های منتخب.'),
p('We average six BA/cost/size values per dataset and use one stability value per method/dataset. Eight datasets, not 240 rows, are analysis blocks [8]. The 2026-10-12 analysis is retrospective: the supplied confirmatory flag and config hash do not establish dated preregistration. We use tie-corrected Friedman tests (df=4), two-sided paired Wilcoxon contrasts of resampled repair versus four controls, and Holm correction within each metric. The supplement also corrects all 24 contrasts and six omnibus tests. Dataset-bootstrap intervals (20,000 draws; seed 20261012) are descriptive, unadjusted percentile intervals.', 'شش مقدار BA، هزینه و اندازه را در هر دیتاست میانگین می‌گیریم و یک مقدار پایداری برای هر روش/دیتاست داریم. هشت دیتاست، نه ۲۴۰ ردیف، بلوک تحلیل‌اند [8]. تحلیل 2026-10-12 گذشته‌نگر است؛ برچسب confirmatory و هش تنظیمات، ثبت پیشینی تاریخ‌دار را ثابت نمی‌کنند. فریدمن با تصحیح تساوی و df=4، ویلکاکسون زوجی دوطرفه برای چهار مقایسه و اصلاح هولم درون هر معیار را به‌کار بردیم. پیوست، اصلاح هم‌زمان ۲۴ مقایسه و شش آزمون کلی را نیز دارد. فاصله‌های صدکی بازنمونه‌گیری دیتاست (۲۰۰۰۰ بار؛ بذر 20261012)، توصیفی و بدون اصلاح چندگانگی‌اند.'),
h('7  Results','۷  نتایج'),
table('Table 5. Equal-weight means over eight eligible datasets; V is final budget exceedance.','جدول ۵. میانگین هم‌وزن هشت دیتاست؛ V تعداد خروجی فراتر از بودجه است.', ['Method','BA','Cost ratio','Size','Stability','V'],mainrows),
table('Table 6. Friedman statistic and raw p; average ranks in E/R/C/M/S order, 1 is best.','جدول ۶. آماره فریدمن و p خام؛ رتبه‌ها به ترتیب E/R/C/M/S، یک بهترین است.', ['Metric','Chi2','p','E','R','C','M','S'],frows),
p('Resampled repair versus EHOA changes cost ratio by -0.19986 (Wilcoxon W=0, raw p=0.00781; within-metric Holm p=0.03125; rank-biserial r=-1). Stability changes by +0.13895 (W=1, raw p=0.01563; Holm p=0.06250; r=0.944). BA changes by +0.00482 (W=9; Holm p=0.750; r=0.500; 95% interval [-0.00149,0.01109]). Neither equivalence nor non-inferiority follows. No resampled repair contrast with the other three repairs is significant within metric for BA or Nogueira stability.', 'تعمیر بازنمونه‌گیری‌شده در برابر EHOA، نسبت هزینه را به‌اندازه -0.19986 تغییر می‌دهد (ویلکاکسون W=0؛ p خام 0.00781؛ هولم درون معیار 0.03125؛ اندازه اثر رتبه‌ای r=-1). تغییر پایداری +0.13895 است (W=1؛ p خام 0.01563؛ هولم 0.06250؛ r=0.944). تغییر BA برابر +0.00482 است (W=9؛ هولم 0.750؛ r=0.500؛ فاصله ۹۵ درصد [-0.00149,0.01109]). این نتایج هم‌ارزی یا عدم‌کاهندگی را ثابت نمی‌کنند. هیچ مقایسه تعمیر بازنمونه‌گیری‌شده با سه تعمیر دیگر، برای BA یا پایداری نوگوئیرا پس از هولم درون معیار معنادار نیست.')]

caserows=[]
for m,label in zip(methods,labels):
    r=next(x for x in bc if x['method']==m)
    caserows.append([label]+[f"{float(r[x]):.4f}" if x!='selected_features' else f"{float(r[x]):.2f}" for x in ['balanced_accuracy','cost_ratio','selected_features','nogueira_stability']])
page4=[h('8  Biomedical case and congress relevance','۸  مطالعه موردی پزشکی و ارتباط با کنگره'),
p('The UCI diagnostic breast-cancer dataset contains 569 fine-needle-aspirate cases (212 malignant, 357 benign) and 30 nuclear-image descriptors [9]. Radius, texture, perimeter, area, smoothness, compactness, concavity, concave points, symmetry and fractal dimension each have mean, error and worst summaries. Table 7 illustrates the five methods on this dataset.', 'دیتاست تشخیصی سرطان پستان UCI شامل ۵۶۹ نمونه آسپیراسیون با سوزن ظریف (۲۱۲ بدخیم و ۳۵۷ خوش‌خیم) و ۳۰ توصیفگر تصویر هسته سلول است [9]. شعاع، بافت، محیط، مساحت، همواری، فشردگی، تقعر، نقاط مقعر، تقارن و بعد فراکتال، هر یک خلاصه میانگین، خطا و بدترین مقدار دارند. جدول ۷ نتیجه پنج روش را در این دیتاست نشان می‌دهد.'),
table('Table 7. Breast-cancer case, means over six runs.','جدول ۷. مطالعه سرطان پستان، میانگین شش اجرا.', ['Method','BA','Cost ratio','Size','Stability'],caserows),
p('Resampled repair selects mean concave points in 6/6 masks and mean texture and mean area in 5/6. These counts describe the archived selections; they do not establish causal biomarkers or clinical reliability. Cheapest repair has greater stability here (0.2815 versus 0.2059). Synthetic feature costs are not biopsy prices; all descriptors derive from the same acquisition, so grouped measurement costs would be more realistic.', 'تعمیر بازنمونه‌گیری‌شده، میانگین نقاط مقعر را در 6/6 و میانگین بافت و مساحت را در 5/6 ماسک انتخاب کرده است. این شمارش، انتخاب‌های آرشیوشده را توصیف می‌کند؛ زیست‌نشانگر علّی یا اتکاپذیری بالینی را ثابت نمی‌کند. تعمیر ارزان‌ترین در این دیتاست پایدارتر است (0.2815 در برابر 0.2059). هزینه مصنوعی ویژگی، قیمت نمونه‌برداری نیست؛ همه توصیفگرها از یک کسب تصویر می‌آیند و هزینه گروهی واقع‌بینانه‌تر خواهد بود.'),
p('The closest official BAIC2026 topic is Artificial Intelligence, Machine Learning and Data Science / Intelligent Disease Diagnosis Systems (baic.um.ac.ir/fa/home/Content/12). Our biomedical contribution concerns computational measurement-panel selection, not a clinically validated diagnostic system. We claim no cognitive-science mechanism.', 'نزدیک‌ترین محور رسمی BAIC2026، «هوش مصنوعی، یادگیری ماشین و علم داده / سیستم‌های هوشمند تشخیص بیماری» است (baic.um.ac.ir/fa/home/Content/12). کاربرد پزشکی پژوهش، انتخاب محاسباتی مجموعه اندازه‌گیری‌هاست و سامانه تشخیصی اعتبارسنجی‌شده بالینی محسوب نمی‌شود. سازوکار علوم شناختی ادعا نمی‌کنیم.'),
h('9  Limitations and reproducibility','۹  محدودیت‌ها و بازتولیدپذیری'),
p('Only one synthetic budget, four MI resamples, eight eligible datasets, two seeds and a short search were assessed. Ordinal Dermatology inputs are treated as continuous by MI; duplicated bootstrap rows can distort nearest-neighbor MI. Overlap-based stability does not verify clinical importance. The excluded datasets need new leakage-free experiments: Steel with 27 predictors and seven classes, Parkinsons with subject-disjoint outer and inner folds. Dated registration, original execution logs and end-to-end timing are unavailable; runtime_seconds excludes MI utility construction in constrained methods.', 'فقط یک بودجه مصنوعی، چهار بازنمونه MI، هشت دیتاست واجد شرایط، دو بذر و جست‌وجوی کوتاه بررسی شد. MI، ورودی‌های ترتیبی Dermatology را پیوسته فرض می‌کند؛ تکرار ردیف‌های bootstrap می‌تواند برآورد همسایگی MI را منحرف کند. پایداری هم‌پوشانی، اهمیت بالینی را تأیید نمی‌کند. دیتاست‌های کنارگذاشته‌شده به آزمایش جدید بدون نشت نیاز دارند: Steel با ۲۷ پیش‌بین و هفت کلاس؛ Parkinsons با تقسیم‌های بیرونی و درونی مستقل در سطح فرد. ثبت تاریخ‌دار، گزارش‌های اجرای اصلی و زمان سرتاسری موجود نیست. runtime_seconds ساخت سودمندی MI در نسخه‌های مقید را شامل نمی‌شود.'),
p('The accompanying local supplement contains all dataset links, 300 mask audits, per-dataset outcomes, cost vectors, statistics and regeneration scripts. Raw files and originals remain intact. No verified public code repository is supplied, and local availability is not public release. The eight-dataset exclusion and analysis are retrospective, not new confirmatory experiments.', 'پیوست محلی، لینک همه داده‌ها، ممیزی ۳۰۰ ماسک، نتایج هر دیتاست، بردار هزینه، آمار و اسکریپت بازتولید را دارد. فایل‌های خام و اصلی حفظ شده‌اند. مخزن عمومی معتبر کد در اختیار نیست؛ دسترسی محلی، انتشار عمومی محسوب نمی‌شود. حذف دو دیتاست و تحلیل هشت‌دیتاستی گذشته‌نگر است و آزمایش تأییدی تازه نیست.'),
h('10  Conclusion','۱۰  نتیجه‌گیری'),
p('All archived constrained outputs are budget-feasible. On eight eligible datasets, resampled repair shows a descriptive cost-stability trade-off but no corrected pairwise stability advantage or general predictive superiority. A stronger biomedical claim requires grouped real costs, subject-safe evaluation and external clinical validation.', 'همه خروجی‌های مقید آرشیوشده شدنی‌اند. در هشت دیتاست واجد شرایط، تعمیر بازنمونه‌گیری‌شده مبادله توصیفی هزینه و پایداری را نشان می‌دهد؛ برتری زوجی پایداری پس از اصلاح یا برتری عمومی پیش‌بینی تأیید نشد. ادعای پزشکی قوی‌تر به هزینه واقعی گروهی، ارزیابی مستقل در سطح فرد و اعتبارسنجی بالینی بیرونی نیاز دارد.'),
h('References','مراجع')]
page1.insert(2,p('Medical feature selection spans imaging, signals and molecular data [10]. Predictive discrimination, reproducible feature identity and acquisition burden are separate objectives. Correlated descriptors can substitute for one another, so a changing mask need not imply poor prediction; conversely, a repeated mask does not establish a valid biomarker. This motivates reporting balanced accuracy, cost and chance-corrected subset stability together.', 'انتخاب ویژگی پزشکی در تصویربرداری، سیگنال و داده مولکولی کاربرد دارد [10]. قدرت تفکیک پیش‌بینی، تکرارپذیری هویت ویژگی و بار کسب اندازه‌گیری، هدف‌های جداگانه‌اند. توصیفگرهای هم‌بسته می‌توانند جای یکدیگر را بگیرند؛ بنابراین تغییر ماسک الزاماً به معنای پیش‌بینی ضعیف نیست و تکرار ماسک نیز اعتبار زیست‌نشانگر را ثابت نمی‌کند. به همین دلیل، دقت متوازن، هزینه و پایداری تصحیح‌شده نسبت به شانس را هم‌زمان گزارش می‌کنیم.'))
page1.insert(-1,p('Classical stability selection aggregates repeated subsample selections and, under specified assumptions and thresholds, controls false selections [11,12]. Our four MI resamples instead estimate a deletion utility before wrapper search. We neither threshold repeated selection probabilities nor inherit those error-control guarantees. Nogueira stability evaluates the final masks [2]; it is not the optimized fitness. Active acquisition [6] chooses measurements sequentially per case, whereas our repair returns a fixed feasible subset.', 'انتخاب پایدار کلاسیک، انتخاب‌های تکرارشده در زیرنمونه‌ها را تجمیع می‌کند و با فرض‌ها و آستانه‌های مشخص، انتخاب کاذب را کنترل می‌کند [11,12]. چهار بازنمونه اطلاعات متقابل در روش حاضر، صرفاً سودمندی حذف را پیش از جست‌وجوی پوششی برآورد می‌کنند. احتمال انتخاب تکرارشده را آستانه‌گذاری نمی‌کنیم و تضمین کنترل خطای آن روش‌ها به این تعمیر منتقل نمی‌شود. شاخص نوگوئیرا ماسک‌های نهایی را ارزیابی می‌کند [2] و تابع برازندگی بهینه‌شده نیست. کسب فعال [6] اندازه‌گیری را برای هر نمونه به‌صورت متوالی انتخاب می‌کند؛ تعمیر حاضر زیرمجموعه ثابت و شدنی برمی‌گرداند.'))
page2.append(p('The continuous-discrete nearest-neighbor MI estimator [13] measures marginal dependence, not conditional contribution after other features are selected. Univariate ranking can retain redundant correlated variables or discard features useful only jointly. With four resamples, the estimated spread is imprecise; lambda=1 expresses a fixed penalty rather than a coverage probability. These design limits motivate comparison with the simpler single-MI and cheapest repairs, rather than attributing every difference to stability.', 'برآوردگر نزدیک‌ترین همسایه اطلاعات متقابل پیوسته–گسسته [13] وابستگی حاشیه‌ای را می‌سنجد، نه سهم شرطی پس از انتخاب سایر ویژگی‌ها. رتبه‌بندی تک‌متغیره ممکن است متغیرهای هم‌بسته زائد را نگه دارد یا ویژگی‌هایی را حذف کند که فقط به‌صورت مشترک مفیدند. با چهار بازنمونه، برآورد پراکندگی دقیق نیست؛ lambda=1 جریمه ثابت است، نه احتمال پوشش. این محدودیت‌ها مقایسه با تعمیر اطلاعات متقابل تک‌نمونه‌ای و ارزان‌ترین را ضروری می‌کنند؛ هر اختلاف را نمی‌توان به پایداری نسبت داد.'))
page3.insert(1,p('SMOTE creates synthetic minority-class training examples [14]; it does not increase the number of independent test subjects. Model selection must remain separate from performance estimation [15]. Here the inner score guides search, while untouched outer tests provide the reported predictions. MI sees the full outer-training labels, including inner-validation labels; thus the inner score is not an unbiased validation estimate. The corrected subject splits address recording overlap, but do not change this utility design.', 'SMOTE نمونه مصنوعی کلاس اقلیت را برای آموزش می‌سازد [14] و تعداد افراد مستقل آزمون را افزایش نمی‌دهد. انتخاب مدل باید از برآورد عملکرد جدا بماند [15]. در این اجراها، امتیاز درونی جست‌وجو را هدایت می‌کند و پیش‌بینی گزارش‌شده روی آزمون بیرونی دست‌نخورده است. اطلاعات متقابل همه برچسب‌های آموزش بیرونی، شامل اعتبارسنجی درونی، را می‌بیند؛ بنابراین امتیاز درونی برآورد اعتبارسنجی بدون سوگیری نیست. تقسیم‌های گروهی اصلاح‌شده، هم‌پوشانی ضبط‌ها را رفع می‌کنند و این طراحی سودمندی را تغییر نمی‌دهند.'))
page1.insert(next(i for i,b in enumerate(page1) if b['kind']=='h' and b['en'].startswith('3 ')),page1.pop(-2))
page2.insert(next(i for i,b in enumerate(page2) if b['kind']=='h' and b['en'].startswith('5 ')),page2.pop())
page2.insert(next(i for i,b in enumerate(page2) if b['kind']=='h' and b['en'].startswith('5 ')),p('Feasibility follows from positive costs and finite deletion: every deletion reduces both subset size and total cost. An unaffordable last feature is cleared; restoration chooses only an individually affordable feature. Thus a nonempty feasible output exists when at least one feature costs at most B, as verified for these datasets. This proves budget compliance, not optimal utility or accuracy. Repair may remove a jointly useful feature because its ranking is marginal and fixed. Repeated repair also maps different search positions to the same mask; cached fitness avoids evaluating such duplicates again.', 'شدنی‌بودن از مثبت‌بودن هزینه‌ها و حذف متناهی نتیجه می‌شود: هر حذف، تعداد ویژگی و هزینه کل را کاهش می‌دهد. آخرین ویژگی غیرقابل‌خرید پاک می‌شود و بازیابی فقط ویژگی منفرد قابل‌خرید را برمی‌گزیند. بنابراین اگر دست‌کم یک ویژگی هزینه‌ای حداکثر برابر B داشته باشد، خروجی ناتهی و شدنی وجود دارد؛ این شرط در داده‌های حاضر بررسی شد. این استدلال رعایت بودجه را ثابت می‌کند، نه سودمندی یا دقت بهینه را. رتبه‌بندی حاشیه‌ای و ثابت ممکن است ویژگی مفید در ترکیب را حذف کند. تعمیر تکراری نیز موقعیت‌های جست‌وجوی متفاوت را به ماسک یکسان نگاشت می‌کند؛ حافظه برازندگی از ارزیابی دوباره این ماسک‌ها جلوگیری می‌کند.'))
page4.insert(next(i for i,b in enumerate(page4) if b['kind']=='h' and b['en'].startswith('9 ')),p('Across the eight datasets, the proposed repair reduces mean synthetic cost by 41.7% and subset size by 26.0% relative to unconstrained EHOA. These comparisons include imposing a budget on an otherwise infeasible baseline: 47 of its 48 outputs exceed that budget. They cannot isolate a benefit of resampling. The cheapest repair already attains mean cost 0.2708 and stability 0.2423, compared with 0.2792 and 0.2550 for resampled repair. The small descriptive stability difference is not corrected pairwise evidence. A medical panel decision therefore needs measured acquisition groups and acceptable sensitivity, rather than choosing the method from stability alone.', 'در هشت دیتاست، تعمیر پیشنهادی میانگین هزینه مصنوعی را 41.7 درصد و اندازه زیرمجموعه را 26.0 درصد نسبت به EHOA نامقید کاهش می‌دهد. این مقایسه شامل اعمال بودجه بر پایه‌ای است که ۴۷ خروجی از ۴۸ خروجی آن از بودجه عبور می‌کنند؛ بنابراین اثر بازنمونه‌گیری را به‌تنهایی جدا نمی‌کند. تعمیر ارزان‌ترین نیز میانگین هزینه 0.2708 و پایداری 0.2423 دارد، در برابر 0.2792 و 0.2550 برای تعمیر بازنمونه‌گیری‌شده. این تفاوت توصیفی کوچک پایداری، شاهد زوجی معنادار پس از اصلاح نیست. تصمیم برای مجموعه اندازه‌گیری پزشکی به گروه‌های کسب واقعی و حساسیت قابل‌قبول نیاز دارد و نباید صرفاً با شاخص پایداری گرفته شود.'))
order=[]
for blocks in [page1,page2,page3,page4]:
    for b in blocks:
        for old in re.findall(r'\[([0-9]+(?:,[0-9]+)*)\]',b.get('en','')+' '+str(b.get('rows',[]))):
            for num in map(int,old.split(',')):
                if num not in order:order.append(num)
mapping={old:new for new,old in enumerate(order,1)}
def renumber(s):
    return re.sub(r'\[([0-9]+(?:,[0-9]+)*)\]',lambda m:'['+','.join(str(mapping[int(x)]) for x in m.group(1).split(','))+']',s)
for blocks in [page1,page2,page3,page4]:
    for b in blocks:
        for lang in ['en','fa']:
            if lang in b:b[lang]=renumber(b[lang])
        if 'rows' in b:b['rows']=[[renumber(x) if isinstance(x,str) else x for x in row] for row in b['rows']]
references=[references[i-1] for i in order]
(OUT/'references_for_submission.txt').write_text('\n'.join(t+' '+u for t,u in references)+'\n',encoding='utf-8')
(OUT/'sources/reference_verification.csv').write_text('reference_id,verified_source,verification_date,doi_status\n'+'\n'.join(f'{i},"{u}",2026-10-12,"'+('Publisher DOI verified' if 'doi.org' in u else 'Direct publisher URL')+'"' for i,(t,u) in enumerate(references,1))+'\n',encoding='utf-8')
for i,(t,u) in enumerate(references,1): page4.append(dict(kind='ref',en=f'[{i}] '+t,fa=f'[{i}] '+t,url=u))
page1=[dict(b,tex=r'\begin{aligned}v_i\leftarrow\operatorname{clip}_{[-12,12]}\big(&w_t v_i+SF_t[2r_1(p_i-x_i)\\&+2r_2(g-x_i)]\big)\end{aligned}',word='vᵢ ← clip[−12,12](wₜvᵢ + SFₜ[2r₁(pᵢ−xᵢ)\n+ 2r₂(g−xᵢ)])') if b['kind']=='eq' and b['tex'].startswith('v_i') else dict(b,tex=r'\begin{gathered}x_i\leftarrow\operatorname{clip}_{[-6,6]}(x_i+v_i)\\w_t=0.9-0.5t/T,\quad SF_t=3-2t/T\end{gathered}',word='xᵢ ← clip[−6,6](xᵢ+vᵢ)\nwₜ = 0.9−0.5t/T; SFₜ = 3−2t/T') if b['kind']=='eq' and b['tex'].startswith('x_i') else b for b in page1]
for b in page2:
    if b['kind']=='table' and len(b['headers'])==6:
        b['widths']=[1.4,1.3,.35,.3,1.9,.65]
        for row in b['rows']:row[4]=row[4].replace('/','/ ')
page3.append(dict(kind='fig',path='figures/eligible_comparison',en='Figure 1. Eight-dataset means; cost uses synthetic units. Stability bars do not imply significant pairwise differences.',fa='شکل ۱. میانگین هشت دیتاست؛ هزینه مصنوعی است. ارتفاع ستون پایداری، معناداری اختلاف زوجی را ثابت نمی‌کند.'))
page4.insert(3,dict(kind='fig',path='figures/breast_top_frequencies',en='Figure 2. Selection counts for the eight most frequent resampled-repair features, from six archived masks.',fa='شکل ۲. شمار انتخاب هشت ویژگی پرتکرار تعمیر بازنمونه‌گیری‌شده در شش ماسک آرشیوشده.'))
if (OUT/'analysis/corrected10_friedman.csv').exists():
    page3.insert(-1,p('As a separate exploratory sensitivity analysis, we completed 60 corrected runs: Steel used 27 predictors and seven classes; Parkinsons used subject-disjoint outer and inner folds after matching all 195 recordings to the UCI file. Combining these with the 240 eligible archived runs gives ten corrected dataset blocks. Friedman BA remains nonsignificant (Chi2=4.32, df=4, p=0.36443). Resampled-versus-EHOA stability has within-metric Holm p=0.01563, but Holm over all 20 contrasts gives p=0.06641. Mean BA delta is -0.00583, with descriptive interval [-0.02523,0.00771]. These new exploratory results neither validate the original leaked experiment nor establish superiority over simpler repairs.', 'در تحلیل حساسیت اکتشافی جداگانه، ۶۰ اجرای اصلاح‌شده انجام دادیم: Steel با ۲۷ پیش‌بین و هفت کلاس؛ Parkinsons با تقسیم بیرونی و درونی مستقل در سطح شناسه فرد، پس از تطبیق تمام ۱۹۵ ضبط با فایل UCI. ترکیب آنها با ۲۴۰ اجرای آرشیوی واجد شرایط، ده بلوک اصلاح‌شده می‌دهد. فریدمن BA همچنان معنادار نیست (Chi2=4.32، df=4، p=0.36443). هولم درون معیار برای اختلاف پایداری پیشنهادی با EHOA برابر 0.01563 است، اما هولم همه ۲۰ مقایسه p=0.06641 می‌دهد. دلتای میانگین BA برابر -0.00583 و فاصله توصیفی [-0.02523,0.00771] است. این نتایج اکتشافی، آزمایش دارای نشت قبلی یا برتری بر تعمیرهای ساده‌تر را تأیید نمی‌کنند.'))
    page4=[dict(b,en=b['en'].replace('The excluded datasets need new leakage-free experiments: Steel with 27 predictors and seven classes, Parkinsons with subject-disjoint outer and inner folds.','The corrected exploratory follow-up is reported separately. UCI describes 31 Parkinson participants, but its file contains 32 distinct subject prefixes; group identity needs independent confirmation.'),fa=b['fa'].replace('دیتاست‌های کنارگذاشته‌شده به آزمایش جدید بدون نشت نیاز دارند: Steel با ۲۷ پیش‌بین و هفت کلاس؛ Parkinsons با تقسیم‌های بیرونی و درونی مستقل در سطح فرد.','اجرای اصلاح‌شده اکتشافی جداگانه گزارش شد. متن UCI تعداد ۳۱ فرد را می‌گوید، ولی فایل ۳۲ پیشوند شناسه دارد؛ هویت گروه‌ها به تأیید مستقل نیاز دارد.')) if b['kind']=='p' else b for b in page4]
for b in page2:
    if b['kind']=='eq' and b['word'].startswith('Φ'):
        b['word']='s²ⱼ = M/(M−1) · pⱼ(1−pⱼ)\nΦ = clip[−1,1](1−mean(s²ⱼ)/[(k̄/D)(1−k̄/D)])'
pages=[page1,page2,page3,page4]
(OUT/'sources/manuscript_content.json').write_text(json.dumps(dict(title=title,abstract=abstract,keywords=keywords,pages=pages),ensure_ascii=False,indent=2),encoding='utf-8')

def esc(s):
    return ''.join({'&':r'\&','%':r'\%','_':r'\_','#':r'\#','{':r'\{','}':r'\}'}.get(c,c) for c in s)
def textext(s,lang):
    if lang=='en': return esc(s)
    return re.sub(r'[-+\[({]*[A-Za-z0-9][A-Za-z0-9 ._/:=+*<>()\[\]{},;-]*',lambda m:r'\lr{'+esc(m.group(0).rstrip())+'}'+(' ' if m.group(0).endswith(' ') else ''),s)

for lang in ['en','fa']:
    base=r'''\documentclass[10pt,a4paper]{article}
\usepackage[top=19mm,bottom=25.4mm,left=16mm,right=16mm]{geometry}
\usepackage{fontspec,graphicx,amsmath,amssymb,multicol,booktabs,array,tabularx}
\usepackage[hidelinks]{hyperref}
\setmainfont{Times New Roman}
\setlength{\columnsep}{6mm}
\setlength{\parindent}{0pt}
\setlength{\parskip}{3pt}
\setlength{\multicolsep}{4pt}
\setlength{\abovedisplayskip}{4pt}
\setlength{\belowdisplayskip}{4pt}
\pagestyle{empty}
\makeatletter
\renewcommand{\section}{\@startsection{section}{1}{0pt}{5pt}{2pt}{\normalfont\normalsize\bfseries}}
\makeatother
\newcommand{\sect}[1]{\section*{#1}}
'''
    if lang=='fa':
        base+=r'''\usepackage{xepersian}
\settextfont{B Nazanin}
\setlatintextfont{Times New Roman}
'''
    base+='\\begin{document}\n'
    for pageindex,blocks in enumerate(pages):
        if pageindex==0:
            base+=r'\begin{center}\includegraphics[width=\textwidth]{baic2026_header.png}\par\vspace{4pt}'+'\n'
            base+='{\\fontsize{14}{17}\\selectfont\\bfseries '+textext(title[lang],lang)+'}\\par\n'
            base+=textext('Amirreza Ghaffarian' if lang=='en' else 'امیررضا غفاریان',lang)+'\\par\n'
            base+=textext('Graduate Student, Islamic Azad University' if lang=='en' else 'دانشجوی کارشناسی ارشد، دانشگاه آزاد اسلامی',lang)+r'\end{center}'+'\n'
            base+=r'\textbf{'+('Abstract. ' if lang=='en' else 'چکیده. ')+r'} '+textext(abstract[lang],lang)+'\n\n'
            base+=r'\textbf{'+('Keywords. ' if lang=='en' else 'کلیدواژه‌ها. ')+r'} '+textext(keywords[lang],lang)+'\n'
        if pageindex==0:base+=r'\begin{multicols}{2}\raggedcolumns'+'\n'
        for block in blocks:
            kind=block['kind']
            if kind=='h': base+=r'\sect{'+textext(block[lang],lang)+'}\n'
            elif kind=='p': base+=textext(block[lang],lang)+'\n\n'
            elif kind=='eq':
                e=r'\begin{equation}\begin{gathered}'+block['tex']+r'\end{gathered}\end{equation}'
                if lang=='fa': e=r'\begin{latin}'+e+r'\end{latin}'
                base+=e+'\n'
            elif kind=='table':
                n=len(block['headers'])
                base+=r'\par\vspace{3pt}\begin{minipage}{\linewidth}{\footnotesize '+textext(block[lang],lang)+r'\par}'+'\n'
                if lang=='fa': base+=r'\begin{latin}'+'\n'
                base+=r'{\fontsize{6.8}{8.3}\selectfont\setlength{\tabcolsep}{2pt}\renewcommand{\arraystretch}{1.1}'+'\n'
                weights=block.get('widths') or [1]*n
                weights=[x*n/sum(weights) for x in weights]
                spec=''.join(r'>{\hsize='+f'{x:.4f}'+r'\hsize\raggedright\arraybackslash}X' for x in weights)
                base+=r'\begin{tabularx}{\linewidth}{'+spec+r'}\toprule'+'\n'
                base+=' & '.join(esc(x) for x in block['headers'])+r'\\\midrule'+'\n'
                for row in block['rows']: base+=' & '.join(esc(str(x)) for x in row)+r'\\'+'\n'
                base+=r'\bottomrule\end{tabularx}}'+'\n'
                if lang=='fa': base+=r'\end{latin}'+'\n'
                base+=r'\end{minipage}\par\vspace{3pt}'+'\n'
            elif kind=='fig':
                base+=r'\par\begin{minipage}{\linewidth}\centering\includegraphics[width=\linewidth]{'+block['path']+r'.pdf}\par{\footnotesize '+textext(block[lang],lang)+r'\par}\end{minipage}\par'+'\n'
            elif kind=='ref':
                t=esc(block[lang])+r' \url{'+block['url']+'}'
                if lang=='fa':t=r'\begin{latin}'+t+r'\end{latin}'
                base+=r'\begin{minipage}{\linewidth}{\fontsize{7}{8.2}\selectfont '+t+r'\par}\end{minipage}\par'+'\n'
        if pageindex==3:base+=r'\end{multicols}'+'\n'
    base+=r'\end{document}'
    (OUT/f'manuscript_{lang}.tex').write_text(base,encoding='utf-8')

def setfont(r,size,fa=False,bold=False):
    r.font.name='Times New Roman';r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=RGBColor(0,0,0)
    rf=r._element.get_or_add_rPr().get_or_add_rFonts()
    for k in ['ascii','hAnsi']:rf.set(qn('w:'+k),'Times New Roman')
    rf.set(qn('w:cs'),'B Nazanin' if fa else 'Times New Roman')
    if fa:
        sz=OxmlElement('w:szCs');sz.set(qn('w:val'),str(round(size*2)));r._element.rPr.append(sz)
def bidi(p,fa):
    x=OxmlElement('w:bidi');x.set(qn('w:val'),'1' if fa else '0');p._p.get_or_add_pPr().append(x)
def addp(doc,text,size,fa=False,bold=False,center=False,style=None):
    pp=doc.add_paragraph(style=style if style and style in doc.styles else 'Normal')
    pp.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    pp.paragraph_format.space_after=Pt(3);pp.paragraph_format.space_before=Pt(0);pp.paragraph_format.line_spacing=1.
    pp.paragraph_format.left_indent=Pt(0);pp.paragraph_format.right_indent=Pt(0);pp.paragraph_format.first_line_indent=Pt(0)
    pp.paragraph_format.widow_control=True
    bidi(pp,fa)
    for s in re.split(r'([-+\[({]*[A-Za-z0-9][A-Za-z0-9 ._/:=+*<>()\[\]{},;-]*)',text) if fa else [text]:
        r=pp.add_run(s);setfont(r,size,fa,bold)
        if fa and re.search('[\u0600-\u06ff]',s):
            rtl=OxmlElement('w:rtl');rtl.set(qn('w:val'),'1');r._element.get_or_add_rPr().append(rtl)
        elif fa:
            rtl=OxmlElement('w:rtl');rtl.set(qn('w:val'),'0');r._element.get_or_add_rPr().append(rtl)
    return pp
def section_settings(s,n,fa=False):
    s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(1.9);s.bottom_margin=Cm(2.54);s.left_margin=s.right_margin=Cm(1.6)
    x=s._sectPr.find(qn('w:cols'))
    if x is not None:s._sectPr.remove(x)
    x=OxmlElement('w:cols');x.set(qn('w:num'),str(n));x.set(qn('w:space'),'340');s._sectPr.append(x)
    old=s._sectPr.find(qn('w:bidi'))
    if old is not None:s._sectPr.remove(old)
    if fa:
        x=OxmlElement('w:bidi');s._sectPr.append(x)
    for part in [s.header,s.footer]:
        for p0 in part.paragraphs:
            p0.clear()
def addtable(doc,b,fa):
    caption=addp(doc,b['fa' if fa else 'en'],8 if not fa else 9,fa,center=True);caption.paragraph_format.keep_with_next=True
    tab=doc.add_table(rows=1,cols=len(b['headers']))
    tab.autofit=False
    weights=b.get('widths') or [1]*len(b['headers'])
    widths=[8.6*x/sum(weights) for x in weights]
    for col,width in zip(tab.columns,widths):col.width=Cm(width)
    tblpr=tab._tbl.tblPr
    tw=tblpr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),'4876')
    old=tblpr.find(qn('w:bidiVisual'))
    if old is not None:tblpr.remove(old)
    align=tblpr.find(qn('w:jc'))
    if align is None:align=OxmlElement('w:jc');tblpr.append(align)
    align.set(qn('w:val'),'left')
    for vals in [b['headers']]+b['rows']:
        row=tab.rows[0] if vals is b['headers'] else tab.add_row()
        for i,val in enumerate(vals):
            cell=row.cells[i];cell.text=''
            cell.width=Cm(widths[i])
            pp=cell.paragraphs[0];pp.alignment=WD_ALIGN_PARAGRAPH.LEFT;pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=1.
            pp.paragraph_format.left_indent=Pt(0);pp.paragraph_format.right_indent=Pt(0);pp.paragraph_format.first_line_indent=Pt(0)
            pp.paragraph_format.keep_with_next=True
            bidi(pp,False)
            r=pp.add_run(str(val));setfont(r,7,False,vals is b['headers'])
            tcpr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders)
            if vals is b['headers']:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E4EBF2');tcpr.append(shade)
            mar=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'25');e.set(qn('w:type'),'dxa');mar.append(e)
            tcpr.append(mar)
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    tab.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

for lang in ['en','fa']:
    fa=lang=='fa';doc=Document(TEMPLATE)
    for child in list(doc._element.body):
        if child.tag!=qn('w:sectPr'):doc._element.body.remove(child)
    for st in doc.styles:
        if st.type==1:
            st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(3)
            pr=st.element.find(qn('w:pPr'))
            if pr is not None:
                num=pr.find(qn('w:numPr'))
                if num is not None:pr.remove(num)
    section_settings(doc.sections[0],1,fa)
    pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.add_run().add_picture(str(OUT/'baic2026_header.png'),width=Cm(17.8))
    addp(doc,title[lang],14,fa,True,True,'Title')
    addp(doc,'امیررضا غفاریان' if fa else 'Amirreza Ghaffarian',10,fa,False,True)
    addp(doc,'دانشجوی کارشناسی ارشد، دانشگاه آزاد اسلامی' if fa else 'Graduate Student, Islamic Azad University',9,fa,False,True)
    addp(doc,('چکیده. ' if fa else 'Abstract. ')+abstract[lang],10 if fa else 9,fa)
    addp(doc,('کلیدواژه‌ها. ' if fa else 'Keywords. ')+keywords[lang],10 if fa else 9,fa)
    section_settings(doc.add_section(WD_SECTION.CONTINUOUS),2,fa)
    for pageindex,blocks in enumerate(pages):
        for b in blocks:
            if b['kind']=='h':
                pp=addp(doc,b[lang],10 if not fa else 11,fa,True,True,'Heading 1');pp.paragraph_format.keep_with_next=True
            elif b['kind']=='p':addp(doc,b[lang],9 if not fa else 10,fa)
            elif b['kind']=='eq':
                pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.space_after=Pt(3)
                math=OxmlElement('m:oMath');arr=OxmlElement('m:eqArr')
                for line in b['word'].split('\n'):
                    ee=OxmlElement('m:e');rr=OxmlElement('m:r');pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'16');pr.append(sz);rr.append(pr);tt=OxmlElement('m:t');tt.text=line;rr.append(tt);ee.append(rr);arr.append(ee)
                math.append(arr);pp._p.append(math)
            elif b['kind']=='table':addtable(doc,b,fa)
            elif b['kind']=='fig':
                pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.add_run().add_picture(str(OUT/(b['path']+'.tiff')),width=Cm(8.6));pp.paragraph_format.keep_with_next=True
                addp(doc,b[lang],8 if not fa else 9,fa,center=True)
            elif b['kind']=='ref':addp(doc,b[lang]+' '+b['url'],7,False)
    section_settings(doc.add_section(WD_SECTION.CONTINUOUS),1,fa)
    doc.core_properties.title=title[lang];doc.core_properties.author='Amirreza Ghaffarian';doc.core_properties.subject='BAIC2026 revised manuscript';doc.core_properties.comments=''
    doc.save(OUT/f'manuscript_{lang}.docx')
print('Wrote synchronized LaTeX and editable Word manuscripts.')
