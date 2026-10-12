from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document

for raw in [r'C:\Users\a.ghafarian\Downloads\EnglishTemp.docx', r'C:\Users\a.ghafarian\Downloads\قالب_مقاله.docx']:
    p = Path(raw)
    print(f'\n=== {p.name} ===')
    with ZipFile(p) as z:
        print('entries', len(z.namelist()))
        print('root rels:', z.read('_rels/.rels').decode('utf-8', errors='replace')[:1200])
        xml = etree.fromstring(z.read('word/document.xml'))
        texts = [''.join(x.itertext()).strip() for x in xml.xpath('.//*[local-name()="p"]')]
        for i, t in enumerate([t for t in texts if t][:120]):
            print(f'{i:03}: {t[:500]}')
    try:
        d = Document(p)
        print('paragraphs', len(d.paragraphs), 'tables', len(d.tables), 'sections', len(d.sections))
        for i, table in enumerate(d.tables[:20]):
            print('TABLE', i, len(table.rows), len(table.columns))
            for row in table.rows[:10]:
                print(' | '.join(c.text.replace('\n', ' / ')[:200] for c in row.cells))
    except Exception as exc:
        print('python-docx error:', repr(exc))
