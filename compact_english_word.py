from docx import Document
from docx.shared import Pt

src = r'C:\Users\a.ghafarian\Downloads\SUBMIT_TO_PROFESSOR_MOATTAR\CostStableEHOA_English_BAIC2026.docx'
out = r'C:\Users\a.ghafarian\Downloads\SUBMIT_TO_PROFESSOR_MOATTAR\CostStableEHOA_English_BAIC2026_Compact.docx'
doc = Document(src)

# PDF-to-Word conversion placed 38 pt after the conference banner.
# Keep the banner and title visually paired without crowding them.
doc.paragraphs[0].paragraph_format.space_after = Pt(6)
doc.paragraphs[1].paragraph_format.space_after = Pt(1)
doc.paragraphs[2].paragraph_format.space_after = Pt(2)
doc.paragraphs[3].paragraph_format.space_after = Pt(0)
doc.paragraphs[4].paragraph_format.space_after = Pt(2)

doc.save(out)
print(out)
